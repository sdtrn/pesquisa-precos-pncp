"""Geração da planilha Excel (.xlsx formatada para folha A4) em fonte Arial com 3 abas e fórmula =AVERAGE(), além de exportação .csv."""

import csv
import html
import math
import re
from io import BytesIO, StringIO
from typing import Any, Dict, List
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


def sanitizar_celula_excel(valor: Any) -> str:
    """
    Remove caracteres de controle invisíveis ilegais na especificação XML do Excel
    (que causavam `IllegalCharacterError` no openpyxl), limpa marcações HTML brutas
    vindas do PNCP e evita interpretação acidental de texto como fórmula quebrada.
    """
    if valor is None:
        return ""
    s = str(valor)
    if not s:
        return ""
    s = html.unescape(s)
    s = re.sub(
        r"</?(?:p|br|strong|b|i|em|u|span|div|ul|ol|li|font|table|tr|td|th)\b[^>]*>",
        " ",
        s,
        flags=re.IGNORECASE,
    )
    s = ILLEGAL_CHARACTERS_RE.sub(" ", s)
    s = re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    if s.startswith(("=", "+", "@")):
        s = "'" + s
    return s


def _calcular_altura_linha_a4(
    textos_e_larguras: List[tuple], altura_minima: float = 24.0, altura_por_linha: float = 13.5
) -> float:
    """Calcula a altura ideal da linha em pontos para caber todo o texto com quebra automática na folha A4."""
    max_linhas = 1
    for texto, larg_col in textos_e_larguras:
        txt_str = str(texto or "")
        if not txt_str or larg_col <= 0:
            continue
        chars_por_linha = max(8, int(larg_col * 0.92))
        linhas_estimadas = math.ceil(len(txt_str) / chars_por_linha)
        if linhas_estimadas > max_linhas:
            max_linhas = linhas_estimadas
    return max(altura_minima, min(180.0, max_linhas * altura_por_linha + 8.0))


def _obter_links_evidencia_item(item: Dict[str, Any]) -> tuple:
    """Retorna `(link_pagina_pncp, link_pdf_principal, resumo_arquivos_str)`."""
    link_pncp = sanitizar_celula_excel(item.get("link_ata_pncp") or "")
    arquivos = item.get("arquivos_arp") or []
    if arquivos and isinstance(arquivos, list):
        primeiro_url = sanitizar_celula_excel(arquivos[0].get("url") or link_pncp)
        partes_txt = []
        for idx_a, arq in enumerate(arquivos, 1):
            tit = sanitizar_celula_excel(arq.get("titulo") or f"Arquivo {idx_a}")
            u = sanitizar_celula_excel(arq.get("url") or "")
            partes_txt.append(f"{tit}: {u}" if u else tit)
        resumo_arqs = " | ".join(partes_txt)
        return link_pncp, primeiro_url, resumo_arqs
    link_dl = sanitizar_celula_excel(item.get("link_download_arp") or link_pncp)
    return link_pncp, link_dl, link_dl


def gerar_planilha_xlsx(
    itens_selecionados: List[Dict[str, Any]],
    itens_janela_top10: List[Dict[str, Any]],
    indices_selecionados: List[int],
    descartados: List[Dict[str, Any]],
) -> bytes:
    """
    Gera o arquivo `.xlsx` em memória (bytes) com fonte Arial em todas as células,
    diagramação ajustada para caber perfeitamente em uma folha A4 e 3 abas:
      1. `Planilha_Parecer`: 6 colunas (incluindo Data de Assinatura) + linha final
         'MÉDIA DOS VALORES UNITÁRIOS' com fórmula `=AVERAGE(B2:B{N})` + Quadro de Links
         Diretos das ARPs Selecionadas (Evidências e Papéis de Trabalho).
      2. `Memoria_de_Pesquisa`: os 10 itens exibidos na janela + metadados completos de auditoria.
      3. `Descartados`: toda ata ou item descartado com o respectivo motivo.
    """
    wb = Workbook()

    # Estilos padronizados em fonte Arial (dimensionados para leitura perfeita em folha A4)
    font_header_a4 = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_body_a4 = Font(name="Arial", size=9, bold=False, color="111111")
    font_bold_a4 = Font(name="Arial", size=9, bold=True, color="1B365D")
    font_link_a4 = Font(name="Arial", size=9, underline="single", color="1B365D")
    font_secao_a4 = Font(name="Arial", size=9, bold=True, color="1B365D")

    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Arial", size=10, bold=False, color="111111")
    font_link = Font(name="Arial", size=10, underline="single", color="1B365D")

    fill_header = PatternFill(
        start_color="1B365D", end_color="1B365D", fill_type="solid"
    )
    fill_sub_header = PatternFill(
        start_color="2C4C7E", end_color="2C4C7E", fill_type="solid"
    )
    fill_media = PatternFill(
        start_color="F4EFE0", end_color="F4EFE0", fill_type="solid"
    )
    fill_zebra = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )

    thin_side = Side(border_style="thin", color="CBD5E1")
    double_bottom = Side(border_style="double", color="B59410")
    top_brass = Side(border_style="medium", color="B59410")

    border_cell = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    border_media = Border(
        left=thin_side, right=thin_side, top=top_brass, bottom=double_bottom
    )

    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=False)

    moeda_format = "R$ #,##0.00"

    # =========================================================================
    # ABA 1: Planilha_Parecer (Diagramada para Folha A4)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Planilha_Parecer"

    # Configuração nativa de impressão em Folha A4 (ajuste automático em 1 página de largura)
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    ws1.page_setup.orientation = ws1.ORIENTATION_PORTRAIT
    ws1.sheet_properties.pageSetUpPr.fitToPage = True
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 0
    ws1.page_setup.horizontalCentered = True
    ws1.page_margins = PageMargins(
        left=0.35, right=0.35, top=0.45, bottom=0.50, header=0.25, footer=0.25
    )
    ws1.oddFooter.left.text = "Pesquisa de Preços em Atas — PNCP (Lei nº 14.133/2021)"
    ws1.oddFooter.right.text = "Desenvolvido por Sérgio Dantas Torres"

    # Larguras proporcionais para caberem com perfeição em folha A4
    larguras_aba1 = [34, 14, 19, 13, 19, 19]
    for i, larg in enumerate(larguras_aba1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = larg

    headers_aba1 = [
        "Descrição do objeto",
        "Valor Unitário (R$)",
        "Nº da ARP / Licitação",
        "Data de Assinatura",
        "Órgão Gerenciador/UF",
        "Fornecedor",
    ]
    ws1.append(headers_aba1)
    ws1.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers_aba1) + 1):
        c = ws1.cell(row=1, column=col_idx)
        c.font = font_header_a4
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_cell

    row_num = 2
    for idx, item in enumerate(itens_selecionados):
        desc_txt = sanitizar_celula_excel(item.get("descricao_objeto"))
        val_unit = float(item.get("valor_unitario") or 0.0)
        arp_txt = sanitizar_celula_excel(item.get("numero_arp_licitacao"))
        dt_ass_txt = sanitizar_celula_excel(item.get("data_assinatura") or "não informado")
        orgao_txt = sanitizar_celula_excel(item.get("orgao_gerenciador_uf"))
        forn_txt = sanitizar_celula_excel(item.get("fornecedor"))
        link_pncp, _, _ = _obter_links_evidencia_item(item)

        ws1.append([desc_txt, val_unit, arp_txt, dt_ass_txt, orgao_txt, forn_txt])
        ws1.row_dimensions[row_num].height = _calcular_altura_linha_a4(
            [
                (desc_txt, larguras_aba1[0]),
                (arp_txt, larguras_aba1[2]),
                (orgao_txt, larguras_aba1[4]),
                (forn_txt, larguras_aba1[5]),
            ]
        )

        for col_idx in range(1, len(headers_aba1) + 1):
            c = ws1.cell(row=row_num, column=col_idx)
            c.font = font_body_a4
            c.border = border_cell
            if idx % 2 == 1:
                c.fill = fill_zebra
            if col_idx == 2:
                c.number_format = moeda_format
                c.alignment = align_right
            elif col_idx == 3:
                c.alignment = align_center
                if link_pncp.startswith("http"):
                    c.hyperlink = link_pncp
                    c.font = font_link_a4
            elif col_idx == 4:
                c.alignment = align_center
            else:
                c.alignment = align_left
        row_num += 1

    # Última linha da tabela principal: MÉDIA DOS VALORES UNITÁRIOS com FÓRMULA (=AVERAGE sobre as linhas selecionadas)
    ultima_linha_dados = max(2, row_num - 1)
    formula_media = (
        f"=AVERAGE(B2:B{ultima_linha_dados})" if itens_selecionados else "=0"
    )
    ws1.append(["MÉDIA DOS VALORES UNITÁRIOS", formula_media, "—", "—", "—", "—"])
    ws1.row_dimensions[row_num].height = 25

    for col_idx in range(1, len(headers_aba1) + 1):
        c = ws1.cell(row=row_num, column=col_idx)
        c.font = font_bold_a4
        c.fill = fill_media
        c.border = border_media
        if col_idx == 2:
            c.number_format = moeda_format
            c.alignment = align_right
        elif col_idx in {3, 4, 5, 6}:
            c.alignment = align_center
        else:
            c.alignment = align_left

    # =========================================================================
    # QUADRO COMPLEMENTAR NA ABA 1: LINKS DIRETOS DAS ARPs SELECIONADAS (EVIDÊNCIAS)
    # =========================================================================
    if itens_selecionados:
        row_ev = row_num + 2
        ws1.merge_cells(start_row=row_ev, start_column=1, end_row=row_ev, end_column=6)
        c_tit = ws1.cell(
            row=row_ev,
            column=1,
            value="EVIDÊNCIAS E PAPÉIS DE TRABALHO — LINKS DIRETOS DE ACESSO E DOWNLOAD DAS ARPs SELECIONADAS",
        )
        c_tit.font = font_secao_a4
        c_tit.fill = fill_media
        c_tit.alignment = align_left
        ws1.row_dimensions[row_ev].height = 22
        for col_i in range(1, 7):
            ws1.cell(row=row_ev, column=col_i).border = border_cell
            ws1.cell(row=row_ev, column=col_i).fill = fill_media

        row_ev += 1
        headers_ev = [
            "Nº da ARP / Licitação",
            "Data de Assinatura",
            "Órgão Gerenciador/UF",
            "Link Direto da ARP no Portal PNCP",
            "Link Direto para Download do PDF/Arquivo da ARP",
            "",
        ]
        ws1.append(headers_ev)
        ws1.merge_cells(start_row=row_ev, start_column=5, end_row=row_ev, end_column=6)
        ws1.row_dimensions[row_ev].height = 22
        for col_i in range(1, 7):
            ce = ws1.cell(row=row_ev, column=col_i)
            ce.font = font_header_a4
            ce.fill = fill_sub_header
            ce.alignment = align_center
            ce.border = border_cell

        for idx_e, item in enumerate(itens_selecionados):
            row_ev += 1
            arp_txt = sanitizar_celula_excel(item.get("numero_arp_licitacao"))
            dt_ass_txt = sanitizar_celula_excel(item.get("data_assinatura") or "não informado")
            orgao_txt = sanitizar_celula_excel(item.get("orgao_gerenciador_uf"))
            link_pncp, link_pdf, _ = _obter_links_evidencia_item(item)

            ws1.append([arp_txt, dt_ass_txt, orgao_txt, link_pncp, link_pdf, ""])
            ws1.merge_cells(start_row=row_ev, start_column=5, end_row=row_ev, end_column=6)
            ws1.row_dimensions[row_ev].height = 22

            for col_i in range(1, 7):
                ce = ws1.cell(row=row_ev, column=col_i)
                ce.font = font_body_a4
                ce.border = border_cell
                if idx_e % 2 == 1:
                    ce.fill = fill_zebra
                if col_i == 2:
                    ce.alignment = align_center
                elif col_i == 4 and link_pncp.startswith("http"):
                    ce.hyperlink = link_pncp
                    ce.font = font_link_a4
                    ce.alignment = align_left
                elif col_i == 5 and link_pdf.startswith("http"):
                    ce.hyperlink = link_pdf
                    ce.font = font_link_a4
                    ce.alignment = align_left
                else:
                    ce.alignment = align_left

    # =========================================================================
    # ABA 2: Memoria_de_Pesquisa
    # =========================================================================
    ws2 = wb.create_sheet(title="Memoria_de_Pesquisa")

    headers_aba2 = [
        "Selecionado (Sim/Não)",
        "Descrição do objeto",
        "Valor Unitário (R$)",
        "Quantitativo previsto na ARP",
        "Unidade",
        "Nº do Item/Lote",
        "Nº da ARP / Licitação",
        "Órgão Gerenciador/UF",
        "Fornecedor",
        "Data de assinatura",
        "Vigência da ata",
        "Data final de vigência",
        "Situação da ata",
        "Esfera",
        "Tipo de órgão",
        "CNPJ do órgão",
        "Link da ata no PNCP",
        "Link Direto Download PDF da ARP",
        "Aderência",
        "Detalhamento da Aderência",
        "Órgão atende ao critério de validade (Sim/Não)",
        "Data/hora da consulta",
        "Escopo dos órgãos escolhido",
        "Status filtrado",
        "Período aplicado",
        "Termos usados na busca",
    ]
    ws2.append(headers_aba2)
    ws2.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers_aba2) + 1):
        c = ws2.cell(row=1, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_cell

    set_indices_sel = set(indices_selecionados)
    for idx, item in enumerate(itens_janela_top10):
        r_idx = idx + 2
        foi_selecionado = "Sim" if idx in set_indices_sel else "Não"
        link_url, link_pdf_url, _ = _obter_links_evidencia_item(item)

        ws2.append(
            [
                foi_selecionado,
                sanitizar_celula_excel(item.get("descricao_objeto")),
                float(item.get("valor_unitario") or 0.0),
                sanitizar_celula_excel(item.get("quantitativo_previsto_arp") or "não informado"),
                sanitizar_celula_excel(item.get("unidade_fornecimento") or "não informado"),
                sanitizar_celula_excel(item.get("numero_item_lote")),
                sanitizar_celula_excel(item.get("numero_arp_licitacao")),
                sanitizar_celula_excel(item.get("orgao_gerenciador_uf")),
                sanitizar_celula_excel(item.get("fornecedor")),
                sanitizar_celula_excel(item.get("data_assinatura") or "não informado"),
                sanitizar_celula_excel(item.get("vigencia_ata") or "não informado"),
                sanitizar_celula_excel(item.get("data_fim_vigencia") or "não informado"),
                sanitizar_celula_excel(item.get("situacao_ata")),
                sanitizar_celula_excel(item.get("esfera")),
                sanitizar_celula_excel(item.get("tipo_orgao")),
                sanitizar_celula_excel(item.get("cnpj_orgao")),
                link_url,
                link_pdf_url,
                sanitizar_celula_excel(item.get("aderencia")),
                sanitizar_celula_excel(item.get("detalhe_aderencia")),
                sanitizar_celula_excel(item.get("orgao_atende_validade")),
                sanitizar_celula_excel(item.get("data_hora_consulta")),
                sanitizar_celula_excel(item.get("escopo_escolhido")),
                sanitizar_celula_excel(item.get("status_escolhido")),
                sanitizar_celula_excel(item.get("periodo_aplicado")),
                sanitizar_celula_excel(item.get("termos_usados")),
            ]
        )
        ws2.row_dimensions[r_idx].height = 36

        for col_idx in range(1, len(headers_aba2) + 1):
            c = ws2.cell(row=r_idx, column=col_idx)
            c.font = font_body
            c.border = border_cell
            if foi_selecionado == "Sim":
                c.fill = fill_media
            elif idx % 2 == 1:
                c.fill = fill_zebra

            if col_idx == 3:
                c.number_format = moeda_format
                c.alignment = align_right
            elif col_idx == 6:
                c.number_format = "@"
                c.alignment = align_center
            elif col_idx == 17 and link_url.startswith("http"):
                c.hyperlink = link_url
                c.font = font_link
                c.alignment = align_left
            elif col_idx == 18 and link_pdf_url.startswith("http"):
                c.hyperlink = link_pdf_url
                c.font = font_link
                c.alignment = align_left
            elif col_idx in {1, 4, 5, 10, 11, 12, 13, 14, 16, 19, 21}:
                c.alignment = align_center
            else:
                c.alignment = align_left

    larguras_aba2 = [
        18, 52, 20, 24, 16, 16, 34, 30, 34, 18, 26, 20, 22, 16, 28, 20, 42, 42, 14, 44, 24, 22, 26, 18, 34, 28
    ]
    for i, larg in enumerate(larguras_aba2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = larg

    # =========================================================================
    # ABA 3: Descartados
    # =========================================================================
    ws3 = wb.create_sheet(title="Descartados")

    headers_aba3 = [
        "Nível (Ata/Item)",
        "Motivo do Descarte",
        "Nº da ARP / Controle PNCP",
        "Órgão Gerenciador/UF",
        "Esfera",
        "Município",
        "Nº do Item/Lote",
        "Descrição do Objeto / Item",
        "Valor Unitário (R$)",
        "Unidade",
        "Data de Assinatura",
        "Vigência da Ata",
        "Situação da Ata",
        "Link no PNCP",
    ]
    ws3.append(headers_aba3)
    ws3.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers_aba3) + 1):
        c = ws3.cell(row=1, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_cell

    for idx, desc in enumerate(descartados):
        r_idx = idx + 2
        val_u = desc.get("Valor Unitário (R$)")
        val_num = None
        if val_u is not None and str(val_u).strip() != "":
            try:
                val_num = float(val_u)
            except Exception:
                val_num = None

        ws3.append(
            [
                sanitizar_celula_excel(desc.get("Nível")),
                sanitizar_celula_excel(desc.get("Motivo do Descarte")),
                sanitizar_celula_excel(desc.get("Nº da ARP / Controle PNCP")),
                sanitizar_celula_excel(desc.get("Órgão Gerenciador/UF")),
                sanitizar_celula_excel(desc.get("Esfera")),
                sanitizar_celula_excel(desc.get("Município")),
                sanitizar_celula_excel(desc.get("Nº do Item/Lote") or "-"),
                sanitizar_celula_excel(desc.get("Descrição do Objeto / Item")),
                val_num if val_num is not None else "-",
                sanitizar_celula_excel(desc.get("Unidade") or "-"),
                sanitizar_celula_excel(desc.get("Data de Assinatura") or "não informado"),
                sanitizar_celula_excel(desc.get("Vigência da Ata") or "não informado"),
                sanitizar_celula_excel(desc.get("Situação da Ata")),
                sanitizar_celula_excel(desc.get("Link no PNCP")),
            ]
        )
        ws3.row_dimensions[r_idx].height = 28

        for col_idx in range(1, len(headers_aba3) + 1):
            c = ws3.cell(row=r_idx, column=col_idx)
            c.font = font_body
            c.border = border_cell
            if col_idx == 7:
                c.number_format = "@"
                c.alignment = align_center
            elif col_idx == 9 and val_num is not None:
                c.number_format = moeda_format
                c.alignment = align_right
            elif col_idx in {1, 5, 6, 10, 11, 12, 13}:
                c.alignment = align_center
            else:
                c.alignment = align_left

    larguras_aba3 = [16, 48, 34, 32, 16, 22, 16, 52, 20, 16, 18, 26, 22, 42]
    for i, larg in enumerate(larguras_aba3, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = larg

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_csv_consolidado(itens_selecionados: List[Dict[str, Any]]) -> bytes:
    """
    Gera o arquivo `.csv` da Tabela Consolidada (codificação UTF-8 com BOM e separador `;`,
    padrão do Excel em português do Brasil), contendo a coluna de Data de Assinatura,
    a média dos valores unitários e os links diretos das ARPs selecionadas.
    """
    out = StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")

    writer.writerow(
        [
            "Descrição do objeto",
            "Valor Unitário (R$)",
            "Nº da ARP / Licitação",
            "Data de Assinatura",
            "Órgão Gerenciador/UF",
            "Fornecedor",
            "Link Direto da ARP no Portal PNCP",
            "Link Direto Download PDF da ARP",
        ]
    )

    soma = 0.0
    for item in itens_selecionados:
        val = float(item.get("valor_unitario") or 0.0)
        soma += val
        val_fmt = f"R$ {val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        link_pncp, link_pdf, _ = _obter_links_evidencia_item(item)
        writer.writerow(
            [
                sanitizar_celula_excel(item.get("descricao_objeto")),
                val_fmt,
                sanitizar_celula_excel(item.get("numero_arp_licitacao")),
                sanitizar_celula_excel(item.get("data_assinatura") or "não informado"),
                sanitizar_celula_excel(item.get("orgao_gerenciador_uf")),
                sanitizar_celula_excel(item.get("fornecedor")),
                link_pncp,
                link_pdf,
            ]
        )

    if itens_selecionados:
        media = soma / len(itens_selecionados)
        media_fmt = f"R$ {media:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        writer.writerow(
            [
                "MÉDIA DOS VALORES UNITÁRIOS",
                media_fmt,
                "—",
                "—",
                "—",
                "—",
                "—",
                "—",
            ]
        )

    return out.getvalue().encode("utf-8-sig")
