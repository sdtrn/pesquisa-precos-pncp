"""Script de execução das 3 rodadas do Teste de Aceite (caso real) com geração de evidências e planilhas .xlsx."""

import json
import sys
from datetime import date
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from src.config import (
    DATA_ASSINATURA,
    ESCOPO_TODOS,
    ESCOPO_VALIDADE,
    STATUS_VIGENTES,
)
from src.exporter import gerar_planilha_xlsx
from src.parser_requisitos import extrair_requisitos_minimos, gerar_termos_busca
from src.pipeline import executar_pipeline_pesquisa, selecionar_top_10_menores


DESCRICAO_TESTE = (
    "Rolo de lona em plástico medindo no mínimo 8MX100M. "
    "Rolo de lona - em plástico resistente, novo de primeiro uso, acondicionado em bobina, "
    "cor preto fosco medindo no mínimo 8mX100m com no mínimo 200 micras (0,200mm). "
    "Pesando no mínimo 110 kg."
)


def executar_rodada(
    num_rodada: int,
    titulo: str,
    escopo: str,
    status: str,
    tipo_data: str,
    dt_ini: date | None,
    dt_fim: date | None,
    termos: list[str],
    requisitos: list[dict],
) -> dict:
    print("\n" + "=" * 80)
    print(f"RODADA {num_rodada}: {titulo}")
    print("=" * 80)

    def prog(msg: str, frac: float) -> None:
        if int(frac * 100) % 25 == 0:
            print(f"  [{frac*100:5.1f}%] {msg}", flush=True)

    res = executar_pipeline_pesquisa(
        descricao_objeto=DESCRICAO_TESTE,
        termos_busca=termos,
        requisitos=requisitos,
        escopo_orgaos=escopo,
        status_ata=status,
        tipo_data_periodo=tipo_data,
        data_inicial=dt_ini,
        data_final=dt_fim,
        progresso_cb=prog,
        usar_cache=True,
    )

    top10_a = selecionar_top_10_menores(res["itens_candidatos"], incluir_revisao_b=False)
    top10_ab = selecionar_top_10_menores(res["itens_candidatos"], incluir_revisao_b=True)

    print(f"\n--- EVIDÊNCIAS DA RODADA {num_rodada} ---")
    print(f"• Escopo: {res['escopo_orgaos']}")
    print(f"• Status: {res['status_ata']}")
    print(f"• Período aplicado: {res['periodo_aplicado_str']}")
    print(f"• Termos de busca usados: {res['termos_busca']}")
    print(f"• Total de atas encontradas no PNCP: {res['total_atas_encontradas']}")
    print(f"• Atas que passaram no filtro de escopo/período: {res['atas_validas_escopo_periodo']}")
    print(f"• Total de itens avaliados nas atas abertas: {res['total_itens_avaliados']}")
    print(f"• Total de registros na aba 'Descartados': {len(res['descartados'])}")
    print(f"• Itens aprovados (1 por ata - R6): A={len(top10_a)} (top 10 A) | A+B={len(top10_ab)} (top 10 A+B)")

    print("\n• Conferência 'Todos = Vigentes + Não vigentes' na API:")
    for c in res["conferencias_status_api"]:
        print(
            f"  - Termo '{c['termo']}': Vigentes ({c['vigentes']}) + Não vigentes ({c['nao_vigentes']}) "
            f"= {c['soma_vigentes_mais_nao_vigentes']} | Todos = {c['todos']} | Válida = {c['equacao_valida']}"
        )
        assert c["equacao_valida"], f"Equação de status falhou para {c['termo']}!"

    print("\n• Descartados por motivo:")
    for mot, qtd in sorted(res["resumo_descartes"].items(), key=lambda x: x[1], reverse=True):
        print(f"  - {mot}: {qtd}")

    print(f"\n• Janela dos 10 menores valores — Modo Padrão (Apenas Aderência A, total={len(top10_a)}):")
    if not top10_a:
        print("  (Nenhum item com todos os 7 atributos explícitos simultaneamente neste filtro)")
    for idx, it in enumerate(top10_a):
        print(
            f"  [{idx+1}] R$ {it['valor_unitario']:,.2f} | Aderência: {it['aderencia']} | "
            f"Qtd ARP: {it['quantitativo_previsto_arp']} | {it['numero_arp_licitacao']} | "
            f"{it['orgao_gerenciador_uf']} | Fornecedor: {it['fornecedor'][:35]}"
        )

    lista_janela = top10_a if len(top10_a) >= 3 else top10_ab
    modo_janela = "Apenas A" if len(top10_a) >= 3 else "Incluindo Revisão Manual B"

    print(f"\n• Janela dos 10 menores valores — {modo_janela} (total={len(lista_janela)}):")
    valores_vistos = []
    for idx, it in enumerate(lista_janela):
        valores_vistos.append(it["valor_unitario"])
        print(
            f"  [{idx+1}] R$ {it['valor_unitario']:,.2f} | Aderência: {it['aderencia']} | "
            f"Qtd ARP: {it['quantitativo_previsto_arp']} | {it['numero_arp_licitacao']} | "
            f"{it['orgao_gerenciador_uf']} | Fornecedor: {it['fornecedor'][:35]} | "
            f"Desc: {it['descricao_objeto'][:70]}..."
        )

    # Verifica ordem estritamente não-decrescente
    assert valores_vistos == sorted(valores_vistos), "Os itens da janela não estão em ordem crescente!"

    # Marca 3 linhas (ou todas se houver menos de 3) e gera a planilha .xlsx
    indices_sel = list(range(min(3, len(lista_janela))))
    itens_sel = [lista_janela[i] for i in indices_sel]

    xlsx_bytes = gerar_planilha_xlsx(
        itens_selecionados=itens_sel,
        itens_janela_top10=lista_janela,
        indices_selecionados=indices_sel,
        descartados=res["descartados"],
    )

    out_path = BASE_DIR / f"evidencia_rodada_{num_rodada}.xlsx"
    with open(out_path, "wb") as f:
        f.write(xlsx_bytes)

    # Inspeciona o .xlsx gerado para validar as 3 abas, fonte Arial e fórmula =AVERAGE
    wb = load_workbook(BytesIO(xlsx_bytes), data_only=False)
    assert wb.sheetnames == ["Planilha_Parecer", "Memoria_de_Pesquisa", "Descartados"], (
        f"Abas incorretas: {wb.sheetnames}"
    )
    ws1 = wb["Planilha_Parecer"]
    linha_media = len(itens_sel) + 2
    rotulo_media = ws1.cell(row=linha_media, column=1).value
    formula_media = ws1.cell(row=linha_media, column=2).value
    fonte_celula = ws1.cell(row=linha_media, column=2).font.name

    print(
        f"\n• Verificação do arquivo '{out_path.name}':"
        f"\n  - Abas: {wb.sheetnames}"
        f"\n  - Linhas selecionadas em Planilha_Parecer: {len(itens_sel)}"
        f"\n  - Linha {linha_media}: '{rotulo_media}' -> Fórmula na coluna B: '{formula_media}' | Fonte: '{fonte_celula}'"
    )
    assert rotulo_media == "MÉDIA DOS VALORES UNITÁRIOS"
    if len(itens_sel) > 0:
        assert formula_media == f"=AVERAGE(B2:B{len(itens_sel)+1})"
    assert fonte_celula == "Arial"

    return res


def main() -> None:
    termos = gerar_termos_busca(DESCRICAO_TESTE)
    requisitos = extrair_requisitos_minimos(DESCRICAO_TESTE)

    print("=" * 80)
    print("PREPARAÇÃO DO TESTE DE ACEITE — TERMOS E REQUISITOS EXTRAÍDOS")
    print("=" * 80)
    print(f"Termos de busca gerados: {termos}")
    print("Requisitos mínimos extraídos:")
    for r in requisitos:
        print(f"  - {r['Atributo']}: {r['Operador']} {r['Valor']} ({r['Unidade']})")

    # Rodada 1: Escopo "Critério de validade", "Vigentes", período vazio
    r1 = executar_rodada(
        num_rodada=1,
        titulo='Escopo "Critério de validade", "Vigentes", período vazio',
        escopo=ESCOPO_VALIDADE,
        status=STATUS_VIGENTES,
        tipo_data=DATA_ASSINATURA,
        dt_ini=None,
        dt_fim=None,
        termos=termos,
        requisitos=requisitos,
    )

    # Rodada 2: Escopo "Critério de validade", "Vigentes", 23/03/2026 a 21/10/2026, "Data de assinatura da ata"
    r2 = executar_rodada(
        num_rodada=2,
        titulo='Escopo "Critério de validade", "Vigentes", Assinatura de 23/03/2026 a 21/10/2026',
        escopo=ESCOPO_VALIDADE,
        status=STATUS_VIGENTES,
        tipo_data=DATA_ASSINATURA,
        dt_ini=date(2026, 3, 23),
        dt_fim=date(2026, 10, 21),
        termos=termos,
        requisitos=requisitos,
    )

    # Rodada 3: Escopo "Todos os órgãos do Brasil", "Vigentes", período vazio
    r3 = executar_rodada(
        num_rodada=3,
        titulo='Escopo "Todos os órgãos do Brasil", "Vigentes", período vazio',
        escopo=ESCOPO_TODOS,
        status=STATUS_VIGENTES,
        tipo_data=DATA_ASSINATURA,
        dt_ini=None,
        dt_fim=None,
        termos=termos,
        requisitos=requisitos,
    )

    print("\n" + "=" * 80)
    print("VERIFICAÇÃO DAS INVARIANTES ENTRE AS 3 RODADAS")
    print("=" * 80)
    print(
        f"1. Atas encontradas na Rodada 2 ({r2['total_atas_encontradas']}) <= Rodada 1 ({r1['total_atas_encontradas']}): "
        f"{r2['total_atas_encontradas'] <= r1['total_atas_encontradas']}"
    )
    print(
        f"2. Atas válidas no escopo/período na Rodada 2 ({r2['atas_validas_escopo_periodo']}) <= Rodada 1 ({r1['atas_validas_escopo_periodo']}): "
        f"{r2['atas_validas_escopo_periodo'] <= r1['atas_validas_escopo_periodo']}"
    )
    assert r2["total_atas_encontradas"] <= r1["total_atas_encontradas"]
    assert r2["atas_validas_escopo_periodo"] <= r1["atas_validas_escopo_periodo"]

    # Verifica que a Rodada 3 contém todas as atas da Rodada 1
    ids_r1 = {
        f"{it['cnpj_orgao']}|{it['numero_arp_licitacao']}"
        for it in r1["itens_candidatos"]
    }
    ids_r3 = {
        f"{it['cnpj_orgao']}|{it['numero_arp_licitacao']}"
        for it in r3["itens_candidatos"]
    }
    r3_contem_r1 = ids_r1.issubset(ids_r3)
    print(
        f"3. Rodada 3 ({r3['atas_validas_escopo_periodo']} atas válidas) contém todas as atas válidas da Rodada 1 ({r1['atas_validas_escopo_periodo']} atas válidas): "
        f"{r3_contem_r1}"
    )
    assert r3_contem_r1, f"Diferença encontrada: {ids_r1 - ids_r3}"

    print("\nTODOS OS TESTES DE ACEITE FORAM CONCLUÍDOS E VALIDADOS COM SUCESSO!")


if __name__ == "__main__":
    main()
